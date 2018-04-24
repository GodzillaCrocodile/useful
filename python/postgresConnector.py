# 'EStroev'
import ipaddress
import psycopg2
from configparser import ConfigParser
import openpyxl


def config(fileName, section):
    parser = ConfigParser()
    parser.read(fileName)

    settings = {}
    if parser.has_section(section):
        params = parser.items(section)
        for param in params:
            settings[param[0]] = param[1]
    else:
        raise Exception('Section {0} not found in the {1} file'.format(section, fileName))

    return settings


def connect(dbSettings):
    try:
        conn = psycopg2.connect(**dbSettings)
        cur = conn.cursor()

        print('[+] Connection to the PostgreSQL established')
        return conn, cur
    except (Exception, psycopg2.DatabaseError) as error:
        print('[-] Error during connection')
        print(error)
        exit(-1)


def create_table(sql, conn, cur):
    try:
        cur.execute(sql)
        conn.commit()
        print('[+] Table created')
    except (Exception, psycopg2.DatabaseError) as error:
        print('[-] Error during creation table')
        print(error)
        exit(-1)


def delete_table(conn, cur, otherSettings):
    deleteTableSQL = 'DROP TABLE %s' % otherSettings['table']
    try:
        cur.execute(deleteTableSQL)
        conn.commit()
        print('[+] Table deleted')
    except (Exception, psycopg2.DatabaseError) as error:
        print('[-] Error during deletion table')
        print(error)
        exit(-1)


def check_table(sql, cur):
    try:
        cur.execute(sql)
        if cur.fetchone()[0]:
            print('[*] Table already exist!')
            return True
    except (Exception, psycopg2.DatabaseError) as error:
        print('[-] Error during check table')
        print(error)
        exit(-1)


def insert(inputFile, table, columns, conn, cur):
    try:
        with open(inputFile, encoding='utf-8') as fileIn:
            cur.copy_from(fileIn, table, columns=columns, sep='|')
            conn.commit()
            print('[+] Data from the %s is written to the %s' % (os.path.basename(inputFile), table))
    except (Exception, psycopg2.DatabaseError) as error:
        print.error('[-] File: %s' % os.path.basename(inputFile))
        exit(-1)


def get_info(sql, cur):
    try:
        cur.execute(sql)
        return cur.fetchall()
    except (Exception, psycopg2.DatabaseError) as error:
        print('[-] Error during get info')
        print(error)
        exit(-1)


def close(conn, cur):
    if cur is not None:
        cur.close()
    if conn is not None:
        conn.close()
    print('[*] Database connection closed')


def write_to_csv(netFile, titles, outPathFile):
    networks = get_network_information(netFile, titles)
    with open(outPathFile, 'w', newline='', encoding='utf-8') as outF:
        csvWriter = csv.writer(outF, delimiter='|')
        for title in networks:
            for network in networks[title]:
                csvWriter.writerow([network, title, networks[title][network]])
            print(f'[+] Write {len(networks[title])} {title} networks entries to {outPathFile}')


def validator(inFile):
    with open(inFile, encoding='utf-8') as inF:
        data = inF.read().split('\n')
    for index, line in enumerate(data, 1):
        if line:
            lines = line.split('|')
            if len(lines) < 3:
                print(f'[-] Invalid line: {index}:"{line}" in {inFile}')


def searcher(ipForSearch):
    try:
        ipaddress.IPv4Network(ipForSearch)
    except ipaddress.AddressValueError:
        print(f'Адрес {ipForSearch} не является верным.')
        exit()
    except ipaddress.NetmaskValueError:
        print(f'Для подсети {ipForSearch} указана слишком большая маска.')
        exit()
    except ValueError:
        print(f'Адрес {ipForSearch} содержит неверный адрес подсети.')
        exit()

    if "/" not in ipForSearch:
        print('Введён IP адрес. Будет осуществлён поиск точного совпадения и всех вхождений.')
        # Соберём подсети для пересчёта, проверяем все подсети в диапазоне [16, 32]
        networks = [ipForSearch + '/' + str(mask) for mask in range(16, 33)]
        flag = 'ip'
    else:
        print('Введена подсеть. Будет осуществлён поиск точного совпадения подсети.')
        networks = []
        networks.append(ipForSearch)
        flag = 'net'

    net_list = []

    if flag == 'ip':
        for net in networks:
            # Рассчитываем подсеть с истинным адресом
            true_adr = str(ipaddress.IPv4Interface(net).network)
            net_list.append(true_adr)

    net_list.append(ipForSearch)

    # Соберём строку запроса
    query = ''
    i = 0
    for adr in net_list:
        if i == 0:
            query = f"SELECT * FROM networks WHERE (network LIKE '{adr}')"
            i += 1
        else:
            query = query + f" OR (network LIKE '{adr}')"

    return query


def xlsx_opener(inFile, title):
    inWB = openpyxl.load_workbook(inFile)
    print(f'[+] Open {inFile}:{title}')
    ws = inWB[title]

    return ws


def load_data(ws, title, networks, cellCount):
    for cell in ws.iter_rows():
        description = list()
        if cell[0].value:
            for i in range(1, cellCount):
                if cell[i].value:
                    description.append(cell[i].value)
            networks[title][cell[0].value] = ';'.join(description)
    print(f'Load {len(networks[title])} network from {title}')

    return networks


def get_network_information(netFile, titles):
    networks = dict()
    for title in titles:
        networks[title] = dict()
        ws = xlsx_opener(netFile, title)
        networks = load_data(ws, title, networks, titles[title])

    return networks


def create_and_write(conn, cur, otherSettings):
    titles = {
        'example': 2
    }

    createTableSQL = '''
    CREATE TABLE %s(
        id SERIAL PRIMARY KEY,
        network TEXT NOT NULL,
        section TEXT NOT NULL,
        description TEXT
    )
    ''' % otherSettings['table']

    checkTableSQL = '''
        SELECT EXISTS (
            SELECT 1
            FROM   information_schema.tables 
            WHERE  table_schema = 'public'
            AND    table_name = '%s'
        )
        ''' % otherSettings['table']

    columns = ('network', 'section', 'description')

    if not check_table(checkTableSQL, cur):
        create_table(createTableSQL, conn, cur)
    write_to_csv(otherSettings['networkfile'], titles, otherSettings['outfile'])
    validator(otherSettings['outfile'])
    insert(otherSettings['outfile'], otherSettings['table'], columns, conn, cur)
    # os.remove(otherSettings['outfile'])
    # print(f"[+] Remove {otherSettings['outfile']}")
    close(conn, cur)


def main():
    settingsPath = 'settings.ini'
    dbSettings = config(settingsPath, 'postgresql')
    otherSettings = config(settingsPath, 'other')

    conn, cur = connect(dbSettings)
    select = "select * from networks where section='PT ABC'"
    ip = ''
    searchQuery = searcher(ip)
    data = get_info(searchQuery, cur)
    # Если что-то найдено - распечатываем результат
    # Если ничего не найдено - сообщаем об этом.
    if len(data) > 0:
        print(f'Сеть {16*" "}Cекция{19*" "}Описание')
        for entry in data:
            print("{1:20}{2:25}{3}".format(*entry))
    else:
        print(f'Для адреса {ip} совпадений в базе не обнаружено.')

    # create_and_write(conn, cur, otherSettings)
    # delete_table(conn, cur, otherSettings)

    # select = "select * from networks where section='CO'"
    # data = get_info(select, cur)
    close(conn, cur)

main()